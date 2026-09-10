# -*- coding: utf-8 -*-
"""审计 GLM-5.3 方案：1) 导出文件是否符合附件5模板；2) 约束细节是否严谨。"""
import sys, os, io, zipfile, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import numpy as np
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
GLM_XLSX = os.path.join(ROOT, 'results', 'glm_original_result1.xlsx')
TPL = r'D:\download\CUMCM2026Problems\C题\附件\附件5\result1.xlsx'

print('=' * 78)
print('A. GLM 导出的 result1.xlsx 结构')
print('=' * 78)
wb = openpyxl.load_workbook(GLM_XLSX)
print('sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print('--- sheet:', ws.title, ws.dimensions)
    for r in range(1, min(ws.max_row, 5) + 1):
        print('   row%-3d %s' % (r, [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]))
    print('   ...')
    for r in range(ws.max_row - 1, ws.max_row + 1):
        print('   row%-3d %s' % (r, [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]))

print()
print('=' * 78)
print('B. 附件5 官方模板结构（对照基准）')
print('=' * 78)
wb2 = openpyxl.load_workbook(TPL)
print('sheets:', wb2.sheetnames)
for ws in wb2.worksheets:
    print('--- sheet:', ws.title, ws.dimensions)
    for r in range(1, 4):
        print('   row%-3d %s' % (r, [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]))
    print('   ... last row:', [ws.cell(ws.max_row, c).value for c in range(1, ws.max_column + 1)])

print()
print('=' * 78)
print('C. 差异判定')
print('=' * 78)
print('工作表名一致      :', wb.sheetnames == wb2.sheetnames)
h1 = [wb['计划购电量'].cell(1, c).value for c in (1, 2)]
h1t = [wb2['计划购电量'].cell(1, c).value for c in (1, 2)]
print('计划购电量 表头    :', h1, ' vs 模板', h1t, '->', h1 == h1t)
print('计划购电量 行数    :', wb['计划购电量'].max_row, ' vs 模板', wb2['计划购电量'].max_row)
h2 = [wb['充放电量'].cell(1, c).value for c in range(1, 6)]
h2t = [wb2['充放电量'].cell(1, c).value for c in range(1, 6)]
print('充放电量 表头      :', h2)
print('              模板 :', h2t, '->', h2 == h2t)
print('充放电量 行数      :', wb['充放电量'].max_row, ' vs 模板', wb2['充放电量'].max_row)
print('时间列类型样例     :', [type(wb['计划购电量'].cell(r, 1).value).__name__ for r in range(2, 8)])

print()
print('=' * 78)
print('D. GLM 解本身的约束细节（d 是否越界、是否等于最大可能弃电量）')
print('=' * 78)
z = np.load(os.path.join(ROOT, 'results', 'glm_original_solution.npz'))
p, c, f, d, E, pi, L, v = (z[k] for k in ('p', 'c', 'f', 'd', 'E', 'pi', 'L', 'v'))
dt = 1 / 6
print('弃电 d 最大值      :', d.max(), ' (光伏 v 最大值 %.2f)' % v.max())
print('d > v 的时段数     :', int(np.sum(d > v + 1e-9)), ' <-- 若>0 说明缺 0<=d<=v 约束')
print('d 总量             : %.4f kWh' % (d.sum() * dt))
print('平衡残差 max       : %.3e kW' % np.max(np.abs(p + v + f - (L + c + d))))
print('同时充放时段数     :', int(np.sum((c > 1e-6) & (f > 1e-6))))
print('购电量 / 购电费    : %.4f kWh / %.4f 元' % (p.sum() * dt, (pi * p * dt).sum()))

# 限定 d<=v 后重解，检验是否改变最优值（用 scipy 复核）
from scipy.optimize import linprog
T = len(pi)
NB, NG, ND, NU, NE = 0, T, 2 * T, 3 * T, 4 * T
nv = 5 * T
cobj = np.zeros(nv); cobj[NB:NB + T] = pi * dt
Aeq, beq = [], []
for t in range(T):   # 平衡：p + v + f = L + c + d  ->  p - c + f - d = L - v
    row = np.zeros(nv); row[NB + t] = 1; row[NG + t] = -1; row[ND + t] = 1; row[NU + t] = -1
    Aeq.append(row); beq.append(L[t] - v[t])
for t in range(T):
    row = np.zeros(nv); row[NE + t] = 1
    if t > 0:
        row[NE + t - 1] = -1
    row[NG + t] = -0.9 * dt; row[ND + t] = dt / 0.9
    Aeq.append(row); beq.append(6000.0 if t == 0 else 0.0)
row = np.zeros(nv); row[NE + T - 1] = 1; Aeq.append(row); beq.append(6000.0)

for tag, ub_d in (('d 无上界（GLM 原式）', None), ('d <= v（增设上界）', v)):
    bounds = [(0, None)] * T + [(0, 5000.0)] * T + [(0, 5000.0)] * T \
             + [(0, (1e9 if ub_d is None else float(ub_d[t]))) for t in range(T)] \
             + [(1200.0, 10800.0)] * T
    r = linprog(cobj, A_eq=np.array(Aeq), b_eq=np.array(beq), bounds=bounds, method='highs')
    print('%-22s success=%s  购电费=%.6f 元  弃电=%.4f kWh' % (
        tag, r.success, r.fun, r.x[NU:NU + T].sum() * dt))
