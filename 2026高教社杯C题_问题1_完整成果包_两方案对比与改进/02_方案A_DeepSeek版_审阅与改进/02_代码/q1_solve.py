# -*- coding: utf-8 -*-
"""
q1_solve.py —— 问题1 求解 + 结果导出（result1.xlsx、CSV、JSON）

运行：
    python code/q1_solve.py

输出：
    output/result1.xlsx           完整计划购电策略（严格套用附件5 模板）
    results/q1_timeseries.csv     144 时段全部决策量与状态量
    results/q1_table1.csv         论文表1（指定时间段购电量/全天购电量/购电费）
    results/q1_table2.csv         论文表2（指定时间段充放电量/0:00与24:00储电量）
    results/q1_summary.json       关键指标汇总
"""
import sys, os, io, json, csv, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import numpy as np
import openpyxl

from common import (load_attachment1, DT, N, RES_DIR, OUT_DIR, FIG_DIR,
                    TPL_RESULT1, TABLE1_CLOCKS, clock_to_t, t_to_interval,
                    summarize_periods_bounds, E_INIT)
from q1_model import solve, build_lp, IDX_G, IDX_D, IDX_U
from scipy.optimize import linprog


def aggregate_tables(res):
    """按论文表1/表2 的口径汇总结果。"""
    buy = res['buy_kWh']

    # ---- 表1：6 个指定 10 分钟区间 ----
    table1 = []
    for c in TABLE1_CLOCKS:
        t = clock_to_t(c)                       # 区间 [c, c+10min) 对应时段 t
        table1.append({
            'interval': '%s-%s' % (c, t_to_interval(t).split('-')[1]),
            't': t,
            'buy_kWh': float(buy[t - 1]),
        })

    # ---- 表2：6 个 4 小时区间 ----
    table2 = []
    for name, a, b in summarize_periods_bounds():
        table2.append({
            'interval': name,
            'charge_kWh': float(res['ch_kWh'][a - 1:b].sum()),
            'discharge_kWh': float(res['dis_kWh'][a - 1:b].sum()),
        })

    return table1, table2


def write_result1_xlsx(res, table2, path):
    """严格按附件5 模板填写 result1.xlsx（保留模板样式）。"""
    shutil.copyfile(TPL_RESULT1, path)
    wb = openpyxl.load_workbook(path)

    ws = wb['计划购电量']
    for i in range(N):
        # 模板第 i+1 行数据（Excel 行号 i+2）对应时段 i+1
        ws.cell(row=i + 2, column=2).value = round(float(res['buy_kWh'][i]), 4)

    ws = wb['充放电量']
    for k, row in enumerate(table2):
        ws.cell(row=k + 2, column=2).value = round(row['charge_kWh'], 4)
        ws.cell(row=k + 2, column=3).value = round(row['discharge_kWh'], 4)
    ws.cell(row=2, column=5).value = round(float(E_INIT), 4)            # 0:00 储电量
    ws.cell(row=3, column=5).value = round(float(res['E_kWh'][-1]), 4)  # 24:00 储电量

    wb.save(path)
    return path


def write_csvs(res, table1, table2):
    # 时间序列
    p = os.path.join(RES_DIR, 'q1_timeseries.csv')
    with open(p, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['t', 'interval', 'price_yuan_per_kWh', 'load_kW', 'pv_kW',
                    'net_load_kW', 'buy_kW', 'buy_kWh', 'ch_kW', 'ch_kWh',
                    'dis_kW', 'dis_kWh', 'curtail_kW', 'E_end_kWh'])
        for i in range(N):
            w.writerow([i + 1, t_to_interval(i + 1),
                        '%.4f' % res['_price'][i], '%.4f' % res['_load'][i], '%.4f' % res['_pv'][i],
                        '%.4f' % (res['_load'][i] - res['_pv'][i]),
                        '%.4f' % res['buy_kW'][i], '%.6f' % res['buy_kWh'][i],
                        '%.4f' % res['ch_kW'][i], '%.6f' % res['ch_kWh'][i],
                        '%.4f' % res['dis_kW'][i], '%.6f' % res['dis_kWh'][i],
                        '%.4f' % res['curtail_kW'][i], '%.4f' % res['E_kWh'][i]])

    # 表1
    p1 = os.path.join(RES_DIR, 'q1_table1.csv')
    with open(p1, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['时间段', '购电量(kWh)'])
        for r in table1:
            w.writerow([r['interval'], '%.4f' % r['buy_kWh']])
        w.writerow(['全天购电量(kWh)', '%.4f' % res['total_buy_kWh']])
        w.writerow(['全天购电费(元)', '%.4f' % res['total_cost']])

    # 表2
    p2 = os.path.join(RES_DIR, 'q1_table2.csv')
    with open(p2, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['时间段', '充电量(kWh)', '放电量(kWh)'])
        for r in table2:
            w.writerow([r['interval'], '%.4f' % r['charge_kWh'], '%.4f' % r['discharge_kWh']])
        w.writerow(['0:00储电量(kWh)', '%.4f' % E_INIT, ''])
        w.writerow(['24:00储电量(kWh)', '%.4f' % res['E_kWh'][-1], ''])
    return p, p1, p2


def baseline_no_storage(price, load, pv):
    """基准方案：储能不动作（充放电功率上界=0），其余同主模型。"""
    c, A_eq, b_eq, bounds = build_lp(price, load, pv, allow_curtail=True)
    bounds[IDX_G:IDX_G + N] = [(0.0, 0.0)] * N
    bounds[IDX_D:IDX_D + N] = [(0.0, 0.0)] * N
    r = linprog(c, A_eq=A_eq, b_eq=b_eq, bounds=bounds, method='highs')
    assert r.success, r.message
    b = r.x[:N]
    return {
        'total_buy_kWh': float(np.sum(b) * DT),
        'total_cost': float(np.sum(price * b) * DT),
        'total_curtail_kWh': float(np.sum(r.x[IDX_U:IDX_U + N]) * DT),
    }


def main():
    df = load_attachment1()
    price = df['price'].to_numpy()
    load = df['load'].to_numpy()
    pv = df['pv'].to_numpy()

    res = solve(price, load, pv, allow_curtail=True, allow_sell=False)
    assert res['success'], res['message']
    res['_price'], res['_load'], res['_pv'] = price, load, pv

    table1, table2 = aggregate_tables(res)
    xlsx_path = write_result1_xlsx(res, table2, os.path.join(OUT_DIR, 'result1.xlsx'))
    write_csvs(res, table1, table2)

    base = baseline_no_storage(price, load, pv)
    summary = {
        'total_buy_kWh': res['total_buy_kWh'],
        'total_cost_yuan': res['total_cost'],
        'avg_price_yuan_per_kWh': res['total_cost'] / res['total_buy_kWh'],
        'total_charge_kWh': res['total_charge_kWh'],
        'total_discharge_kWh': res['total_discharge_kWh'],
        'total_curtail_kWh': res['total_curtail_kWh'],
        'pv_total_kWh': res['pv_total_kWh'],
        'pv_used_kWh': res['pv_used_kWh'],
        'pv_utilization': res['pv_used_kWh'] / res['pv_total_kWh'],
        'load_total_kWh': res['load_total_kWh'],
        'net_load_total_kWh': float(np.sum(load - pv) * DT),
        'storage_loss_kWh': res['total_charge_kWh'] - res['total_discharge_kWh'],
        'E_min': float(res['E_kWh'].min()),
        'E_max': float(res['E_kWh'].max()),
        'E_0': float(E_INIT),
        'E_T': float(res['E_kWh'][-1]),
        'baseline_no_storage': base,
        'saving_yuan': base['total_cost'] - res['total_cost'],
        'saving_ratio': 1.0 - res['total_cost'] / base['total_cost'],
        'table1': table1,
        'table2': table2,
    }
    with open(os.path.join(RES_DIR, 'q1_summary.json'), 'w', encoding='utf-8') as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print('求解状态:', res['message'])
    print('全天购电量 = %.4f kWh' % summary['total_buy_kWh'])
    print('全天购电费 = %.4f 元  (平均电价 %.4f 元/kWh)' % (summary['total_cost_yuan'], summary['avg_price_yuan_per_kWh']))
    print('储能充电总量 = %.4f kWh, 放电总量 = %.4f kWh, 循环损耗 = %.4f kWh'
          % (summary['total_charge_kWh'], summary['total_discharge_kWh'], summary['storage_loss_kWh']))
    print('光伏发电总量 = %.4f kWh, 实际消纳 = %.4f kWh (利用率 %.2f%%), 弃光 = %.4f kWh'
          % (summary['pv_total_kWh'], summary['pv_used_kWh'], summary['pv_utilization'] * 100, summary['total_curtail_kWh']))
    print('储电量区间 [%.2f, %.2f]，0:00 = %.2f，24:00 = %.2f' % (summary['E_min'], summary['E_max'], summary['E_0'], summary['E_T']))
    print('对照（无储能）：购电费 %.4f 元，购电量 %.4f kWh，弃光 %.4f kWh' % (
        base['total_cost'], base['total_buy_kWh'], base['total_curtail_kWh']))
    print('储能套利节约 = %.4f 元 (%.2f%%)' % (summary['saving_yuan'], summary['saving_ratio'] * 100))
    print('已写出:', xlsx_path)


if __name__ == '__main__':
    main()
